from __future__ import annotations

from dataclasses import dataclass
from copy import deepcopy
import json
from pathlib import Path
import re
from typing import Any
from uuid import UUID

from docx import Document
from PIL import Image
from app.ai.brand_intelligence import BrandIntelligenceService
from app.ai.providers.base import PromptEnvelope
from app.ai.providers.router import ProviderRouter
from app.ai.rag.ocr import OCRService
from app.ai.tone_intelligence import ToneIntelligenceService
from app.core.enums import BrandSpaceLifecycle
from app.core.exceptions import GenerationFailureError, LifecycleError, NotFoundError
from app.integrations.object_storage import LocalObjectStorage
from app.models.content import ContentSession, ContentVersion
from app.repositories.brand import BrandSpaceRepository, ObjectiveRepository, PersonaRepository
from app.repositories.content import AssetRepository, ContentRepository
from app.repositories.knowledge import KnowledgeAssetRepository
from app.schemas.common import StudioPanelSelection
from app.services.artifact_state import ArtifactStateService
from app.services.content import ContentService
from app.services.content_planning import ContentPlanningService
from app.services.format_family_planning import FormatFamilyPlanningService
from app.services.live_research import LiveResearchService
from app.services.research_editorial_planning import ResearchEditorialPlanningService


@dataclass(slots=True)
class TextGenerationResult:
    content_version: ContentVersion
    assistant_payload: dict[str, Any]
    assistant_text: str


class TextContentService:
    DELIVERABLE_CONTRACTS: dict[str, dict[str, Any]] = {
        "blog": {
            "label": "blog article",
            "body_instruction": "Return a complete markdown article with a strong title, introduction, clear section headings, and a concise conclusion.",
            "headline_instruction": "Headline should be a publishable article title.",
            "cta_instruction": "CTA should be optional and subtle unless the prompt clearly asks for one.",
            "hashtags_instruction": "Do not use hashtags for blogs unless explicitly requested.",
            "metadata_fields": ["outline", "seo_keywords", "sources_used"],
        },
        "linkedin_post": {
            "label": "LinkedIn post",
            "body_instruction": "Return a polished LinkedIn post with a strong hook, short paragraphs, and a thoughtful closing insight.",
            "headline_instruction": "Headline should be short and optional; if used, it must feel native to LinkedIn.",
            "cta_instruction": "CTA should feel conversational, not salesy.",
            "hashtags_instruction": "Use only a small set of relevant hashtags if they genuinely help.",
            "metadata_fields": ["hook_type", "key_takeaway", "sources_used"],
        },
        "instagram_caption": {
            "label": "Instagram caption",
            "body_instruction": "Return a short, high-clarity caption with strong first-line pull and clean rhythm.",
            "headline_instruction": "Headline is optional and should stay very short if present.",
            "cta_instruction": "CTA should feel natural and lightweight.",
            "hashtags_instruction": "Include a tight set of relevant hashtags if useful.",
            "metadata_fields": ["hook_type", "caption_style", "sources_used"],
        },
        "x_post": {
            "label": "X post",
            "body_instruction": "Return a concise, sharp single post with one clear insight and no filler.",
            "headline_instruction": "Headline is optional and should be very short if used.",
            "cta_instruction": "CTA should be minimal.",
            "hashtags_instruction": "Avoid hashtags unless explicitly useful.",
            "metadata_fields": ["hook_type", "key_takeaway", "sources_used"],
        },
        "x_thread": {
            "label": "X thread",
            "body_instruction": "Return a numbered thread in plain text with 4-8 compact posts, each adding a fresh point.",
            "headline_instruction": "Headline should frame the thread, not repeat the first post.",
            "cta_instruction": "CTA should only appear in the final post if needed.",
            "hashtags_instruction": "Avoid hashtags unless explicitly requested.",
            "metadata_fields": ["thread_count", "thread_structure", "sources_used"],
        },
        "youtube_description": {
            "label": "YouTube description",
            "body_instruction": "Return a complete YouTube description with a strong opening summary, scannable sections, and helpful context.",
            "headline_instruction": "Headline should support discoverability without clickbait.",
            "cta_instruction": "CTA can include subscribe/watch/learn prompts if relevant.",
            "hashtags_instruction": "Hashtags are optional and should be limited.",
            "metadata_fields": ["seo_keywords", "chapters", "sources_used"],
        },
        "newsletter": {
            "label": "newsletter draft",
            "body_instruction": "Return a polished newsletter body with clear sections and a smooth editorial flow.",
            "headline_instruction": "Headline should feel editorial and relevant.",
            "cta_instruction": "CTA should be useful and outcome-led.",
            "hashtags_instruction": "Do not use hashtags.",
            "metadata_fields": ["outline", "key_takeaway", "sources_used"],
        },
        "email": {
            "label": "email draft",
            "body_instruction": "Return a complete email body with a clear opening, concise middle, and direct close.",
            "headline_instruction": "Headline should work as a subject-line style opener if appropriate.",
            "cta_instruction": "CTA should be direct and specific.",
            "hashtags_instruction": "Do not use hashtags.",
            "metadata_fields": ["subject_line", "preview_text", "sources_used"],
        },
        "script": {
            "label": "script draft",
            "body_instruction": "Return a complete script with clear spoken flow and natural cadence.",
            "headline_instruction": "Headline should frame the script topic cleanly.",
            "cta_instruction": "CTA is optional and only if the prompt warrants it.",
            "hashtags_instruction": "Do not use hashtags.",
            "metadata_fields": ["outline", "estimated_length", "sources_used"],
        },
        "general_copy": {
            "label": "content draft",
            "body_instruction": "Return a polished, complete draft tailored to the prompt and platform cues.",
            "headline_instruction": "Headline should be useful, clear, and native to the format.",
            "cta_instruction": "CTA should match the prompt and stay natural.",
            "hashtags_instruction": "Use hashtags only if they clearly fit the format.",
            "metadata_fields": ["key_takeaway", "sources_used"],
        },
        "social_caption": {
            "label": "social caption",
            "body_instruction": "Return a concise social caption with a clear opening and a natural close.",
            "headline_instruction": "Headline is optional and should stay brief.",
            "cta_instruction": "CTA should stay light and platform-native.",
            "hashtags_instruction": "Use only a small relevant set of hashtags if needed.",
            "metadata_fields": ["hook_type", "key_takeaway", "sources_used"],
        },
        "long_description": {
            "label": "description draft",
            "body_instruction": "Return a complete description with clear structure and practical detail.",
            "headline_instruction": "Headline should clearly orient the reader.",
            "cta_instruction": "CTA is optional unless requested.",
            "hashtags_instruction": "Only include hashtags if relevant to the requested format.",
            "metadata_fields": ["outline", "sources_used"],
        },
    }

    def __init__(self, session) -> None:
        self.session = session
        self.providers = ProviderRouter()
        self.brands = BrandSpaceRepository(session)
        self.personas = PersonaRepository(session)
        self.objectives = ObjectiveRepository(session)
        self.contents = ContentRepository(session)
        self.assets = AssetRepository(session)
        self.knowledge_assets = KnowledgeAssetRepository(session)
        self.content = ContentService(session)
        self.live_research = LiveResearchService()
        self.research_editorial = ResearchEditorialPlanningService()
        self.format_family_planning = FormatFamilyPlanningService()
        self.content_planning = ContentPlanningService()
        self.artifacts = ArtifactStateService()
        self.storage = LocalObjectStorage()
        self.ocr = OCRService()
        self.tone = ToneIntelligenceService()

    async def generate(
        self,
        *,
        tenant_id: UUID,
        brand_space_id: UUID,
        user_id: UUID,
        session: ContentSession,
        prompt: str,
        studio_panel: StudioPanelSelection,
        persona_id: UUID | None = None,
        objective_id: UUID | None = None,
        deliverable_type: str | None = None,
        uses_previous_output: bool = False,
    ) -> TextGenerationResult:
        artifact_service = getattr(self, "artifacts", ArtifactStateService())
        brand = await self.brands.get_scoped(tenant_id, brand_space_id)
        if not brand:
            raise NotFoundError("Brand Space not found")
        if str(getattr(brand, "lifecycle_state", BrandSpaceLifecycle.ACTIVE)) != BrandSpaceLifecycle.ACTIVE:
            raise LifecycleError("Brand Space must be Active for chat")

        session_memory = await self.content._build_session_memory(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            session=session,
            current_prompt=prompt,
        )
        personas = await self.personas.list_by_brand(brand_space_id, tenant_id)
        objectives = await self.objectives.list_by_brand(brand_space_id, tenant_id)
        selected_persona = next((item for item in personas if item.id == persona_id), None)
        selected_objective = next((item for item in objectives if item.id == objective_id), None)
        brand_context = dict(getattr(brand, "resolved_brand_context", {}) or {})
        persona_context = BrandIntelligenceService.persona_to_dict(selected_persona)
        objective_context = BrandIntelligenceService.objective_to_dict(selected_objective)

        retrieved_knowledge, knowledge_state = await self.content._build_retrieved_knowledge(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            prompt=prompt,
            studio_panel=studio_panel.model_dump(),
        )
        knowledge_brief = self._knowledge_brief(retrieved_knowledge)
        live_research = self.live_research.gather_sync(
            prompt=prompt,
            studio_panel=studio_panel.model_dump(),
            compiled_context={"knowledge_brief": knowledge_brief},
        )
        planning_bundle = self.content_planning.build_text_plan(
            prompt=prompt,
            studio_panel=studio_panel.model_dump(),
            brand_context=brand_context,
            persona_context=persona_context,
            objective_context=objective_context,
            knowledge_brief=knowledge_brief,
            live_research=live_research,
            deliverable_type=deliverable_type or "general_copy",
        )
        research_editorial_brief = planning_bundle["research_editorial_brief"]
        self._assert_research_guard(prompt=prompt, brief=research_editorial_brief, stage="text.generate")
        format_family_plan = planning_bundle["format_family_plan"]
        content_plan = planning_bundle["content_plan"]
        generated_payload = self._generate_payload(
            prompt=prompt,
            brand_context=brand_context,
            persona_context=persona_context,
            objective_context=objective_context,
            deliverable_type=deliverable_type or "general_copy",
            session_memory=session_memory,
            knowledge_brief=knowledge_brief,
            live_research=live_research,
            research_editorial_brief=research_editorial_brief,
            format_family_plan=format_family_plan,
            content_plan=content_plan,
            uses_previous_output=uses_previous_output,
            previous_output=str(session.conversational_context.get("last_text_output") or "").strip() or None,
            revision_scope=None,
        )
        content_text = self._content_text(generated_payload)
        tone_feedback = self.tone.evaluate(
            content=content_text,
            brand_context=brand_context,
            persona_context=persona_context,
            content_payload=generated_payload,
            objective_context=objective_context,
            message_strategy={},
        )
        artifact_state = artifact_service.build_content_state(
            mode="content_only",
            prompt=prompt,
            studio_panel=studio_panel.model_dump(),
            research_objects={
                "retrieval_channels": list(retrieved_knowledge.keys()),
                "knowledge_state": knowledge_state,
                "live_research": live_research,
                "research_editorial_brief": research_editorial_brief,
            },
            planning_objects={
                "deliverable_type": deliverable_type or "general_copy",
                "format_family_plan": format_family_plan,
                "content_plan": content_plan,
            },
            revision_lineage=artifact_service.build_revision_lineage(
                parent_version_id=self._parse_uuid_or_none(session.conversational_context.get("last_content_version_id")),
                rewrite_mode="fresh_generation",
            ),
        )

        content_version = ContentVersion(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            session_id=session.id,
            parent_version_id=self._parse_uuid_or_none(session.conversational_context.get("last_content_version_id")),
            created_by=user_id,
            lifecycle_state="generated",
            content_type="text_content",
            title=str(generated_payload.get("headline") or generated_payload.get("title") or "").strip() or None,
            prompt=prompt,
            selected_persona_id=persona_id,
            selected_template_id=None,
            objective_id=objective_id,
            studio_panel=studio_panel.model_dump(),
            generated_payload=generated_payload,
            blueprint_payload={},
            explainability_metadata={
                "mode": "content_only",
                "deliverable_type": deliverable_type or "general_copy",
                "brand_context_snapshot": brand_context,
                "selected_persona": persona_context,
                "selected_objective": objective_context,
                "session_memory": session_memory,
                "retrieved_knowledge": retrieved_knowledge,
                "retrieval_channels": list(retrieved_knowledge.keys()),
                "knowledge_state": knowledge_state,
                "live_research": live_research,
                "research_editorial_brief": research_editorial_brief,
                "format_family_plan": format_family_plan,
                "content_plan": content_plan,
                "artifact_state": artifact_state,
            },
            tone_score=int(tone_feedback.get("score") or 0),
            tone_feedback=tone_feedback,
        )
        await self.contents.add(content_version)
        await self.session.flush()

        assistant_payload = {
            "mode": "content_only",
            "deliverable_type": deliverable_type or "general_copy",
            "content_version_id": str(content_version.id),
            "generated_payload": generated_payload,
            "tone_feedback": tone_feedback,
            "live_research": live_research,
            "research_editorial_brief": research_editorial_brief,
            "format_family_plan": format_family_plan,
            "content_plan": content_plan,
            "retrieval_channels": list(retrieved_knowledge.keys()),
            "artifact_state": artifact_state,
        }
        assistant_text = self._assistant_text(generated_payload)
        return TextGenerationResult(
            content_version=content_version,
            assistant_payload=assistant_payload,
            assistant_text=assistant_text,
        )

    async def evaluate(
        self,
        *,
        tenant_id: UUID,
        brand_space_id: UUID,
        session: ContentSession,
        prompt: str,
        persona_id: UUID | None = None,
        objective_id: UUID | None = None,
        reference_asset_ids: list[UUID] | None = None,
    ) -> dict[str, Any]:
        from app.services.evaluation import EvaluationService

        evaluation_service = EvaluationService(getattr(self, "session", None))
        evaluation_service.helper = self
        for attribute in ("brands", "personas", "objectives", "tone", "artifacts", "providers"):
            if hasattr(self, attribute):
                setattr(evaluation_service, attribute, getattr(self, attribute))
        return await evaluation_service.evaluate(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            session=session,
            prompt=prompt,
            persona_id=persona_id,
            objective_id=objective_id,
            reference_asset_ids=reference_asset_ids,
        )

    @staticmethod
    def _assert_research_guard(*, prompt: str, brief: dict[str, Any], stage: str) -> None:
        guard = brief.get("research_guard") if isinstance(brief.get("research_guard"), dict) else {}
        if not guard or not bool(guard.get("hard_fail")):
            return
        raise GenerationFailureError(
            str(guard.get("reason") or "Research-backed generation requirements were not met."),
            failure_type="missing_research",
            reason_code="research_backing_required",
            user_safe_message=(
                "I couldn't generate this safely because the prompt needs externally verified research, "
                "but live source verification was unavailable. Please try again or attach a source document."
            ),
            retryable=True,
            rule_source="system",
            suggested_next_action="Retry with live research available or upload a supporting source.",
            details={"stage": stage, "prompt_excerpt": str(prompt or "")[:180]},
        )

    async def rewrite(
        self,
        *,
        tenant_id: UUID,
        brand_space_id: UUID,
        user_id: UUID,
        session: ContentSession,
        content_version_id: UUID,
        rewrite_instruction: str,
        studio_panel: StudioPanelSelection,
        revision_scope: dict[str, Any] | None = None,
    ) -> TextGenerationResult:
        artifact_service = getattr(self, "artifacts", ArtifactStateService())
        original = await self.contents.get_scoped(content_version_id, tenant_id, brand_space_id)
        if not original:
            raise NotFoundError("Content version not found")
        if str(getattr(original, "content_type", "") or "") != "text_content":
            raise NotFoundError("Content version is not a text content draft")

        explainability = original.explainability_metadata if isinstance(original.explainability_metadata, dict) else {}
        deliverable_type = str(
            explainability.get("deliverable_type")
            or (original.generated_payload or {}).get("deliverable_type")
            or session.conversational_context.get("last_text_deliverable_type")
            or "general_copy"
        ).strip() or "general_copy"
        brand_context = dict(
            getattr(await self.brands.get_scoped(tenant_id, brand_space_id), "resolved_brand_context", {}) or {}
        ) or dict(explainability.get("brand_context_snapshot") or {})
        persona_context = dict(explainability.get("selected_persona") or {})
        objective_context = dict(explainability.get("selected_objective") or {})
        session_memory = await self.content._build_session_memory(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            session=session,
            current_prompt=rewrite_instruction,
        )
        knowledge_brief = self._knowledge_brief(explainability.get("retrieved_knowledge", {}))
        live_research = explainability.get("live_research", {}) if isinstance(explainability.get("live_research"), dict) else {}
        planning_bundle = self.content_planning.build_text_plan(
            prompt=rewrite_instruction,
            studio_panel=studio_panel.model_dump(),
            brand_context=brand_context,
            persona_context=persona_context,
            objective_context=objective_context,
            knowledge_brief=knowledge_brief,
            live_research=live_research,
            deliverable_type=deliverable_type,
        )
        research_editorial_brief = planning_bundle["research_editorial_brief"]
        self._assert_research_guard(prompt=rewrite_instruction, brief=research_editorial_brief, stage="text.rewrite")
        format_family_plan = planning_bundle["format_family_plan"]
        content_plan = planning_bundle["content_plan"]
        original_payload = original.generated_payload if isinstance(original.generated_payload, dict) else {}
        generated_payload = self._generate_payload(
            prompt=rewrite_instruction,
            brand_context=brand_context,
            persona_context=persona_context,
            objective_context=objective_context,
            deliverable_type=deliverable_type,
            session_memory=session_memory,
            knowledge_brief=knowledge_brief,
            live_research=live_research,
            research_editorial_brief=research_editorial_brief,
            format_family_plan=format_family_plan,
            content_plan=content_plan,
            uses_previous_output=True,
            previous_output=self._assistant_text(original_payload),
            revision_scope=revision_scope,
        )
        generated_payload = self._apply_revision_scope_to_payload(
            original_payload=original_payload,
            rewritten_payload=generated_payload,
            revision_scope=revision_scope,
        )
        content_text = self._content_text(generated_payload)
        tone_feedback = self.tone.evaluate(
            content=content_text,
            brand_context=brand_context,
            persona_context=persona_context,
            content_payload=generated_payload,
            objective_context=objective_context,
            message_strategy={},
        )
        prior_artifact_state = explainability.get("artifact_state") if isinstance(explainability.get("artifact_state"), dict) else {}
        artifact_state = artifact_service.build_content_state(
            mode="content_only",
            prompt=rewrite_instruction,
            studio_panel=studio_panel.model_dump(),
            research_objects={
                "retrieval_channels": list((explainability.get("retrieved_knowledge") or {}).keys()),
                "knowledge_state": explainability.get("knowledge_state"),
                "live_research": live_research,
                "research_editorial_brief": research_editorial_brief,
            },
            planning_objects={
                "deliverable_type": deliverable_type,
                "format_family_plan": format_family_plan,
                "content_plan": content_plan,
            },
            revision_lineage=artifact_service.build_revision_lineage(
                parent_version_id=original.id,
                source_content_version_id=original.id,
                rewrite_mode="text_revision",
                rewrite_instruction=rewrite_instruction,
                revision_scope=revision_scope,
                prior_lineage=prior_artifact_state.get("revision_lineage"),
            ),
            source_linked_artifacts={
                "source_content_version_id": str(original.id),
            },
        )
        content_version = ContentVersion(
            tenant_id=tenant_id,
            brand_space_id=brand_space_id,
            session_id=session.id,
            parent_version_id=original.id,
            created_by=user_id,
            lifecycle_state="edited",
            content_type="text_content",
            title=str(generated_payload.get("headline") or generated_payload.get("title") or original.title or "").strip() or None,
            prompt=rewrite_instruction,
            selected_persona_id=original.selected_persona_id,
            selected_template_id=None,
            objective_id=original.objective_id,
            studio_panel=studio_panel.model_dump(),
            generated_payload=generated_payload,
            blueprint_payload={},
            explainability_metadata={
                **deepcopy(explainability),
                "mode": "content_only",
                "deliverable_type": deliverable_type,
                "brand_context_snapshot": brand_context,
                "selected_persona": persona_context,
                "selected_objective": objective_context,
                "session_memory": session_memory,
                "live_research": live_research,
                "research_editorial_brief": research_editorial_brief,
                "format_family_plan": format_family_plan,
                "content_plan": content_plan,
                "rewrite_mode": "text_revision",
                "rewrite_instruction": rewrite_instruction,
                "rewrite_source_content_version_id": str(original.id),
                "revision_scope": revision_scope or {},
                "artifact_state": artifact_state,
            },
            tone_score=int(tone_feedback.get("score") or 0),
            tone_feedback=tone_feedback,
        )
        await self.contents.add(content_version)
        await self.session.flush()
        assistant_payload = {
            "mode": "content_only",
            "deliverable_type": deliverable_type,
            "content_version_id": str(content_version.id),
            "generated_payload": generated_payload,
            "tone_feedback": tone_feedback,
            "live_research": live_research,
            "research_editorial_brief": research_editorial_brief,
            "format_family_plan": format_family_plan,
            "content_plan": content_plan,
            "retrieval_channels": list((explainability.get("retrieved_knowledge") or {}).keys()),
            "rewrite_mode": "text_revision",
            "rewrite_source_content_version_id": str(original.id),
            "revision_scope": revision_scope or {},
            "artifact_state": artifact_state,
        }
        return TextGenerationResult(
            content_version=content_version,
            assistant_payload=assistant_payload,
            assistant_text=self._assistant_text(generated_payload),
        )

    def _generate_payload(
        self,
        *,
        prompt: str,
        brand_context: dict[str, Any],
        persona_context: dict[str, Any],
        objective_context: dict[str, Any],
        deliverable_type: str,
        session_memory: dict[str, Any],
        knowledge_brief: list[dict[str, Any]],
        live_research: dict[str, Any],
        research_editorial_brief: dict[str, Any],
        format_family_plan: dict[str, Any],
        content_plan: dict[str, Any],
        uses_previous_output: bool,
        previous_output: str | None,
        revision_scope: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        provider = self.providers.get_text_provider("generation")
        contract = self._deliverable_contract(deliverable_type)
        prompt_safe_research_brief = self._prompt_safe_research_editorial_brief(research_editorial_brief)
        prompt_safe_format_plan = self._prompt_safe_plan(format_family_plan)
        prompt_safe_content_plan = self._prompt_safe_plan(content_plan)
        fallback = self._fallback_payload(
            prompt=prompt,
            deliverable_type=deliverable_type,
            previous_output=previous_output,
        )
        payload = provider.generate_structured_json(
            PromptEnvelope(
                system=(
                    "You are a research-oriented branded content strategist. "
                    "Generate polished text-only deliverables that follow the brand voice, audience, and objective context. "
                    "Return JSON only with keys: deliverable_type, headline, body, cta, hashtags, metadata. "
                    "The body must be fully usable as the final content. "
                    "If live research or retrieved knowledge is present, use it to make the writing concrete and non-generic."
                ),
                user=(
                    f"Prompt: {prompt}\n"
                    f"Deliverable type: {deliverable_type}\n"
                    f"Deliverable contract: {contract}\n"
                    f"Brand context: {brand_context}\n"
                    f"Persona context: {persona_context}\n"
                    f"Objective context: {objective_context}\n"
                    f"Session memory: {session_memory}\n"
                    f"Retrieved knowledge: {knowledge_brief}\n"
                    f"Live research: {live_research}\n"
                    f"Research editorial brief: {prompt_safe_research_brief}\n"
                    f"Format family plan: {prompt_safe_format_plan}\n"
                    f"Content plan: {prompt_safe_content_plan}\n"
                    f"Previous output: {previous_output or ''}\n"
                    f"Uses previous output: {uses_previous_output}\n"
                    f"Revision scope: {revision_scope or {}}\n"
                    "If Uses previous output is true, revise the prior draft instead of starting over and preserve the original intent unless the new instruction changes it.\n"
                    f"{self._revision_scope_instruction(revision_scope)}\n"
                    "If Research editorial brief is active, preserve its thesis, angle, reader payoff, hook strategy, and outline. "
                    "Do not flatten research-heavy prompts into generic summaries or generic brand filler.\n"
                    "When verified facts are present, treat them as the only claims that can be stated as confirmed facts. "
                    "Use inferences as interpretation, not as hard fact, and surface uncertainties when the evidence is incomplete or conditional.\n"
                    "Follow the citation discipline implied by the research editorial brief and format family plan so exact facts, dates, and numbers stay source-backed without making the copy clunky.\n"
                    "Treat Format family plan as the authoritative planning contract for this content family. "
                    "Follow its primary unit, body shape, content structure, metadata fields, and density guidance instead of using one generic structure for every deliverable.\n"
                    "Treat Content plan as the authoritative narrative plan for this draft. "
                    "If it specifies sections, ordered beats, or slide-by-slide intent, follow that structure explicitly and keep each unit meaningfully distinct.\n"
                    "Populate metadata only with useful fields for this deliverable. "
                    "Do not mention being an AI. Do not output placeholders."
                ),
            ),
            fallback=fallback,
        )
        normalized = self._normalize_payload(payload, fallback=fallback, deliverable_type=deliverable_type)
        self._assert_not_model_fallback(
            prompt=prompt,
            normalized_payload=normalized,
            deliverable_type=deliverable_type,
            previous_output=previous_output,
            stage="text.generate_payload",
        )
        return self.research_editorial.enforce_source_backing(
            normalized,
            prompt_text=prompt,
            brief=research_editorial_brief,
        )

    @staticmethod
    def _prompt_safe_research_editorial_brief(value: dict[str, Any]) -> dict[str, Any]:
        brief = dict(value or {})
        outline: list[dict[str, Any]] = []
        for item in (brief.get("outline") or [])[:8]:
            if not isinstance(item, dict):
                continue
            outline.append(
                {
                    "index": str(item.get("index") or "").strip(),
                    "role": str(item.get("role") or "").strip(),
                    "purpose": str(item.get("purpose") or "").strip(),
                }
            )
        brief["outline"] = outline
        return brief

    @staticmethod
    def _prompt_safe_plan(value: dict[str, Any]) -> dict[str, Any]:
        plan = dict(value or {})
        plan.pop("notes", None)
        return plan

    @staticmethod
    def _revision_scope_instruction(revision_scope: dict[str, Any] | None) -> str:
        if not isinstance(revision_scope, dict) or not revision_scope:
            return "No granular revision scope is active."
        instructions: list[str] = []
        targeted_fields = [
            str(field).strip()
            for field in revision_scope.get("targeted_fields", [])
            if str(field).strip()
        ]
        slide_indexes = [
            int(value)
            for value in revision_scope.get("slide_indexes", [])
            if str(value).strip().isdigit()
        ]
        slide_targets = [
            str(target).strip()
            for target in revision_scope.get("slide_targets", [])
            if str(target).strip()
        ]
        if targeted_fields:
            instructions.append(f"Only rewrite these targeted fields unless the instruction clearly requires more: {', '.join(targeted_fields)}.")
        if slide_indexes or slide_targets:
            slide_scope = [*(f"slide {index}" for index in slide_indexes), *slide_targets]
            instructions.append(f"Keep the requested edits scoped to {' and '.join(slide_scope)} where possible.")
        if revision_scope.get("preserve_visuals"):
            instructions.append("Preserve the existing visual framing and avoid changing design-oriented direction in the copy.")
        if revision_scope.get("preserve_copy"):
            instructions.append("Preserve the existing copy/message as much as possible and only support the requested visual/layout change.")
        if revision_scope.get("change_layout"):
            instructions.append("The request implies a layout/composition change; avoid rewriting unaffected narrative content.")
        if revision_scope.get("change_tone"):
            instructions.append("Apply the requested tone shift while preserving the core facts, argument, and deliverable structure.")
        if revision_scope.get("only_targeted"):
            instructions.append("Do not broaden the rewrite beyond the explicitly requested scope.")
        return " ".join(instructions) if instructions else "No granular revision scope is active."

    @staticmethod
    def _fallback_payload(*, prompt: str, deliverable_type: str, previous_output: str | None) -> dict[str, Any]:
        headline = prompt.strip().splitlines()[0][:120].strip() or "Content Draft"
        body = previous_output or ""
        return {
            "deliverable_type": deliverable_type,
            "headline": headline,
            "body": body,
            "cta": "",
            "hashtags": [],
            "metadata": {"source": "fallback", "sources_used": []},
        }

    @staticmethod
    def _assert_not_model_fallback(
        *,
        prompt: str,
        normalized_payload: dict[str, Any],
        deliverable_type: str,
        previous_output: str | None,
        stage: str,
    ) -> None:
        metadata = normalized_payload.get("metadata") if isinstance(normalized_payload.get("metadata"), dict) else {}
        if str(metadata.get("source") or "").strip().lower() != "fallback":
            return
        raise GenerationFailureError(
            "Structured text generation fell back to a non-authoritative placeholder payload.",
            failure_type="provider_failure",
            reason_code="text_generation_fallback",
            user_safe_message=(
                "I couldn't produce a strong final draft because the text model fell back instead of returning a usable result. "
                "Please retry the generation."
            ),
            retryable=True,
            rule_source="system",
            suggested_next_action="Retry the generation.",
            details={
                "stage": stage,
                "deliverable_type": deliverable_type,
                "prompt_excerpt": str(prompt or "")[:180],
                "had_previous_output": bool(str(previous_output or "").strip()),
            },
        )

    @staticmethod
    def _normalize_payload(payload: dict[str, Any], *, fallback: dict[str, Any], deliverable_type: str) -> dict[str, Any]:
        if not isinstance(payload, dict):
            payload = {}
        metadata = payload.get("metadata") if isinstance(payload.get("metadata"), dict) else {}
        hashtags = payload.get("hashtags") if isinstance(payload.get("hashtags"), list) else fallback.get("hashtags", [])
        return {
            "deliverable_type": str(payload.get("deliverable_type") or deliverable_type).strip() or deliverable_type,
            "headline": str(payload.get("headline") or fallback.get("headline") or "").strip(),
            "body": str(payload.get("body") or fallback.get("body") or "").strip(),
            "cta": str(payload.get("cta") or fallback.get("cta") or "").strip(),
            "hashtags": [str(item).strip() for item in hashtags if str(item).strip()],
            "metadata": metadata,
        }

    @staticmethod
    def _apply_revision_scope_to_payload(
        *,
        original_payload: dict[str, Any],
        rewritten_payload: dict[str, Any],
        revision_scope: dict[str, Any] | None,
    ) -> dict[str, Any]:
        if not isinstance(rewritten_payload, dict):
            return rewritten_payload
        if not isinstance(revision_scope, dict) or not revision_scope:
            return rewritten_payload

        targeted_fields = {
            str(field).strip()
            for field in revision_scope.get("targeted_fields", [])
            if str(field).strip()
        }
        only_targeted = bool(revision_scope.get("only_targeted"))
        preserve_copy = bool(revision_scope.get("preserve_copy"))
        has_explicit_scope = bool(targeted_fields) or only_targeted or preserve_copy
        if not has_explicit_scope:
            return rewritten_payload

        merged_payload = deepcopy(rewritten_payload)
        for field in ("headline", "body", "cta", "hashtags", "metadata"):
            if field in targeted_fields and not preserve_copy:
                continue
            if field == "metadata":
                original_metadata = original_payload.get("metadata")
                if isinstance(original_metadata, dict):
                    merged_payload["metadata"] = deepcopy(original_metadata)
                continue
            if field == "hashtags":
                merged_payload["hashtags"] = list(original_payload.get("hashtags", []) or [])
                continue
            merged_payload[field] = str(original_payload.get(field) or "").strip()
        if original_payload.get("deliverable_type"):
            merged_payload["deliverable_type"] = str(original_payload.get("deliverable_type") or "").strip()
        return merged_payload

    @classmethod
    def _deliverable_contract(cls, deliverable_type: str) -> dict[str, Any]:
        return dict(cls.DELIVERABLE_CONTRACTS.get(deliverable_type, cls.DELIVERABLE_CONTRACTS["general_copy"]))

    @staticmethod
    def _assistant_text(payload: dict[str, Any]) -> str:
        parts = [
            str(payload.get("headline") or "").strip(),
            str(payload.get("body") or "").strip(),
            str(payload.get("cta") or "").strip(),
        ]
        hashtags = payload.get("hashtags") if isinstance(payload.get("hashtags"), list) else []
        if hashtags:
            parts.append(" ".join(str(tag).strip() for tag in hashtags if str(tag).strip()))
        return "\n\n".join(part for part in parts if part)

    @staticmethod
    def _content_text(payload: dict[str, Any]) -> str:
        return "\n\n".join(
            part
            for part in [
                str(payload.get("headline") or "").strip(),
                str(payload.get("body") or "").strip(),
                str(payload.get("cta") or "").strip(),
            ]
            if part
        ).strip()

    @staticmethod
    def _knowledge_brief(retrieved_knowledge: dict[str, list[dict]]) -> list[dict[str, Any]]:
        brief: list[dict[str, Any]] = []
        for channel, items in (retrieved_knowledge or {}).items():
            for item in items[:2]:
                if not isinstance(item, dict):
                    continue
                brief.append(
                    {
                        "channel": channel,
                        "content": str(item.get("content") or "").strip()[:600],
                        "source_url": str(item.get("source_url") or "").strip() or None,
                    }
                )
        return brief[:8]

    @staticmethod
    def _parse_uuid_or_none(value: Any) -> UUID | None:
        try:
            return UUID(str(value))
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _evaluation_target_text(
        *,
        prompt: str,
        session: ContentSession,
        asset_review_blocks: list[dict[str, str]] | None = None,
    ) -> str:
        text = str(prompt or "").strip()
        for fence in ('"""', "```"):
            if fence in text:
                parts = text.split(fence)
                if len(parts) >= 3:
                    return parts[1].strip()
        if ":" in text and any(marker in text.casefold() for marker in ("tone", "brand", "guideline", "consistency", "compliance")):
            _, candidate = text.split(":", 1)
            if candidate.strip():
                return candidate.strip()
        if asset_review_blocks:
            return "\n\n".join(
                block["content"].strip()
                for block in asset_review_blocks
                if str(block.get("content") or "").strip()
            ).strip()
        return str(session.conversational_context.get("last_text_output") or "").strip()

    async def _review_text_from_assets(
        self,
        *,
        tenant_id: UUID,
        brand_space_id: UUID,
        reference_asset_ids: list[UUID],
        review_prompt: str,
        brand_context: dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        blocks: list[dict[str, Any]] = []
        for asset_id in reference_asset_ids:
            generated_asset = await self.assets.get_scoped(asset_id, tenant_id, brand_space_id)
            if generated_asset:
                content_version = await self.contents.get_scoped(generated_asset.content_version_id, tenant_id, brand_space_id)
                if content_version:
                    payload = content_version.generated_payload if isinstance(content_version.generated_payload, dict) else {}
                    content = self._assistant_text(payload).strip()
                    visual_review = self._visual_review_from_storage_path(
                        storage_path=str(getattr(generated_asset, "storage_path", "") or "").strip(),
                        mime_type=str(getattr(generated_asset, "mime_type", "") or "").strip(),
                        asset_kind=self._asset_kind_from_mime(str(getattr(generated_asset, "mime_type", "") or "").strip()),
                        expected_prompt=str(getattr(content_version, "prompt", "") or review_prompt or "").strip(),
                        brand_context=brand_context or {},
                    )
                    if content or visual_review:
                        content = content or str(visual_review.get("ocr_text") or "").strip()
                        blocks.append(
                            {
                                "asset_id": str(asset_id),
                                "source_type": "generated_asset",
                                "asset_name": str(getattr(content_version, "title", "") or "Generated content").strip() or "Generated content",
                                "mime_type": str(getattr(generated_asset, "mime_type", "") or "").strip() or None,
                                "asset_kind": self._asset_kind_from_mime(str(getattr(generated_asset, "mime_type", "") or "").strip()),
                                "page_count": int(((generated_asset.metadata_json or {}).get("slide_count") or 1)),
                            "content_preview": content[:240],
                            "extraction_method": "generated_payload",
                            "review_workflow": "visual_first" if visual_review else "text_first",
                            "diagnostics": {
                                "dimensions": {
                                    "width": getattr(generated_asset, "width", None),
                                    "height": getattr(generated_asset, "height", None),
                                },
                                    "render_source": (generated_asset.metadata_json or {}).get("render_source"),
                                    "slide_index": (generated_asset.metadata_json or {}).get("slide_index"),
                                    "slide_count": (generated_asset.metadata_json or {}).get("slide_count"),
                                    **(visual_review.get("diagnostics", {}) if isinstance(visual_review, dict) else {}),
                                },
                                "visual_review": visual_review,
                                "content": content,
                            }
                        )
                continue

            knowledge_asset = await self.knowledge_assets.get_scoped(asset_id, tenant_id, brand_space_id)
            if knowledge_asset:
                detail = self._knowledge_asset_review_detail(knowledge_asset)
                content = str(detail.get("content") or "").strip()
                visual_review = self._visual_review_from_storage_path(
                    storage_path=str(getattr(knowledge_asset, "storage_path", "") or "").strip(),
                    mime_type=str(getattr(knowledge_asset, "mime_type", "") or "").strip(),
                    asset_kind=self._knowledge_asset_kind(knowledge_asset),
                    expected_prompt=str(review_prompt or "").strip(),
                    brand_context=brand_context or {},
                )
                if not content and visual_review:
                    content = str(visual_review.get("ocr_text") or "").strip()
                if content or visual_review:
                    blocks.append(
                        {
                            "asset_id": str(asset_id),
                            "source_type": "knowledge_asset",
                            "asset_name": str(getattr(knowledge_asset, "name", "") or getattr(knowledge_asset, "original_filename", "") or "Uploaded asset").strip() or "Uploaded asset",
                            "mime_type": str(getattr(knowledge_asset, "mime_type", "") or "").strip() or None,
                            "asset_kind": self._knowledge_asset_kind(knowledge_asset),
                            "page_count": int(getattr(knowledge_asset, "page_count", 0) or 0),
                            "field_key": str(getattr(knowledge_asset, "field_key", "") or "").strip() or None,
                            "content_preview": content[:240],
                            "extraction_method": detail.get("extraction_method"),
                            "review_workflow": "visual_first" if visual_review else "text_first",
                            "gap_note": detail.get("gap_note"),
                            "diagnostics": {
                                **self._knowledge_asset_diagnostics(knowledge_asset),
                                **(visual_review.get("diagnostics", {}) if isinstance(visual_review, dict) else {}),
                            },
                            "visual_review": visual_review,
                            "content": content,
                        }
                    )
        return blocks

    def _visual_review_from_storage_path(
        self,
        *,
        storage_path: str,
        mime_type: str,
        asset_kind: str,
        expected_prompt: str,
        brand_context: dict[str, Any],
    ) -> dict[str, Any]:
        if not storage_path or asset_kind not in {"image", "document", "presentation"}:
            return {}
        storage = getattr(self, "storage", None)
        ocr = getattr(self, "ocr", None)
        if storage is None or ocr is None:
            return {}
        if not storage.exists(storage_path):
            return {}
        absolute_path = storage.absolute_path(storage_path)
        try:
            extracted = ocr.extract(absolute_path)
        except Exception as exc:  # noqa: BLE001
            return {
                "warnings": [f"Visual review could not run OCR/analysis for this asset: {exc}"],
                "diagnostics": {
                    "prompt_alignment_score": 0,
                    "layout_readability_score": 0,
                    "density_score": 0,
                    "brand_alignment_score": 0,
                    "visual_diagnostic_score": 0,
                    "warning_count": 1,
                },
                "page_reviews": [],
                "ocr_text": "",
                "summary": "Visual review could not inspect this asset because OCR/analysis failed.",
                "findings": ["OCR/analysis was unavailable for this asset during evaluation."],
            }

        ocr_text = str(extracted.get("text") or "").strip()
        warnings = [str(item).strip() for item in (extracted.get("warnings") or []) if str(item).strip()]
        source_format = str(extracted.get("source_format") or "").strip()
        page_count = int(extracted.get("page_count") or 0)
        page_images = [str(path).strip() for path in (extracted.get("images") or []) if str(path).strip()]
        if not page_images and asset_kind == "image":
            page_images = [absolute_path]
        page_reviews: list[dict[str, Any]] = []
        if page_images:
            for index, image_path in enumerate(page_images, start=1):
                page_text = ocr_text if len(page_images) == 1 else ""
                analysis: dict[str, Any] = {}
                if len(page_images) > 1:
                    try:
                        page_extracted = ocr.extract(image_path)
                    except Exception:  # noqa: BLE001
                        page_extracted = {}
                    page_text = str(page_extracted.get("text") or "").strip()
                    analysis = self._read_visual_analysis_path(str(page_extracted.get("analysis_path") or "").strip())
                if not analysis:
                    analysis = self._read_visual_analysis_for_image_path(image_path)
                page_reviews.append(
                    self._build_visual_page_review(
                        page_index=index,
                        image_path=image_path,
                        analysis=analysis,
                        expected_prompt=expected_prompt,
                        page_text=page_text,
                        brand_context=brand_context,
                    )
                )
        else:
            page_reviews.append(
                self._build_visual_page_review(
                    page_index=1,
                    image_path=absolute_path,
                    analysis=self._read_visual_analysis_path(str(extracted.get("analysis_path") or "").strip()),
                    expected_prompt=expected_prompt,
                    page_text=ocr_text,
                    brand_context=brand_context,
                )
            )

        prompt_alignment_score = self._average_visual_metric(page_reviews, "prompt_alignment_score")
        layout_readability_score = self._average_visual_metric(page_reviews, "layout_readability_score")
        density_score = self._average_visual_metric(page_reviews, "density_score")
        brand_alignment_score = self._average_visual_metric(page_reviews, "brand_alignment_score")
        hierarchy_score = self._average_visual_metric(page_reviews, "hierarchy_score")
        crowding_score = self._average_visual_metric(page_reviews, "crowding_score")
        page_balance_score = self._average_visual_metric(page_reviews, "page_balance_score")
        ocr_confidence_score = self._average_visual_metric(page_reviews, "ocr_confidence_score")
        visual_diagnostic_score = int(
            round(
                (
                    prompt_alignment_score
                    + layout_readability_score
                    + density_score
                    + brand_alignment_score
                    + hierarchy_score
                    + crowding_score
                    + page_balance_score
                    + ocr_confidence_score
                )
                / 8.0
            )
        )
        findings = self._visual_findings_from_pages(page_reviews, expected_prompt=expected_prompt)
        document_segments = self._document_segments_from_page_reviews(page_reviews, asset_kind=asset_kind)
        region_overview = self._region_overview_from_pages(page_reviews)
        summary = (
            f"Visual review across {len(page_reviews)} page(s): prompt alignment {prompt_alignment_score}/100, "
            f"readability {layout_readability_score}/100, density {density_score}/100, brand alignment {brand_alignment_score}/100, "
            f"hierarchy {hierarchy_score}/100, OCR confidence {ocr_confidence_score}/100."
        )
        return {
            "review_mode": "visual_first",
            "source_format": source_format or Path(absolute_path).suffix.lower().lstrip("."),
            "page_count": page_count or len(page_reviews),
            "ocr_text": ocr_text[:4000],
            "warnings": warnings[:8],
            "document_segments": document_segments,
            "region_overview": region_overview,
            "page_reviews": page_reviews,
            "findings": findings,
            "summary": summary,
            "diagnostics": {
                "prompt_alignment_score": prompt_alignment_score,
                "layout_readability_score": layout_readability_score,
                "density_score": density_score,
                "brand_alignment_score": brand_alignment_score,
                "hierarchy_score": hierarchy_score,
                "crowding_score": crowding_score,
                "page_balance_score": page_balance_score,
                "ocr_confidence_score": ocr_confidence_score,
                "visual_diagnostic_score": visual_diagnostic_score,
                "warning_count": len(warnings),
            },
        }

    @staticmethod
    def _read_visual_analysis_path(analysis_path: str) -> dict[str, Any]:
        if not analysis_path:
            return {}
        path = Path(analysis_path)
        if not path.exists():
            return {}
        try:
            parsed = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        except Exception:  # noqa: BLE001
            return {}
        return parsed if isinstance(parsed, dict) else {}

    def _read_visual_analysis_for_image_path(self, image_path: str) -> dict[str, Any]:
        image = Path(image_path)
        direct = image.with_name(f"{image.stem}_analysis.json")
        analysis = self._read_visual_analysis_path(str(direct))
        if analysis:
            return analysis
        ocr = getattr(self, "ocr", None)
        if ocr is None:
            return {}
        try:
            extracted = ocr.extract(image_path)
        except Exception:  # noqa: BLE001
            return {}
        return self._read_visual_analysis_path(str(extracted.get("analysis_path") or "").strip())

    @staticmethod
    def _prompt_topic_tokens(prompt: str, *, limit: int = 12) -> list[str]:
        stopwords = set(ContentService.TOPIC_STOPWORDS) | {
            "about",
            "against",
            "check",
            "consistency",
            "evaluate",
            "guideline",
            "guidelines",
            "image",
            "review",
            "score",
            "tone",
            "visual",
        }
        tokens: list[str] = []
        seen: set[str] = set()
        for token in re.findall(r"[A-Za-z][A-Za-z0-9'-]{2,}", str(prompt or "").lower()):
            if token in stopwords or token in seen:
                continue
            seen.add(token)
            tokens.append(token)
            if len(tokens) >= limit:
                break
        return tokens

    @staticmethod
    def _extract_brand_name_tokens(brand_context: dict[str, Any]) -> list[str]:
        tokens: list[str] = []
        for candidate in (
            brand_context.get("brand_name"),
            (brand_context.get("identity") or {}).get("brand_name") if isinstance(brand_context.get("identity"), dict) else None,
            (brand_context.get("identity") or {}).get("name") if isinstance(brand_context.get("identity"), dict) else None,
        ):
            for token in re.findall(r"[A-Za-z][A-Za-z0-9'-]{2,}", str(candidate or "").lower()):
                if token not in tokens:
                    tokens.append(token)
        return tokens

    @staticmethod
    def _extract_brand_palette_hexes(brand_context: dict[str, Any]) -> list[str]:
        results: list[str] = []
        seen: set[str] = set()

        def visit(value: Any) -> None:
            if isinstance(value, dict):
                for key, nested in value.items():
                    if key in {"hex", "hex_code"}:
                        hex_code = str(nested or "").strip().upper()
                        if hex_code and not hex_code.startswith("#") and re.fullmatch(r"[0-9A-F]{6}", hex_code):
                            hex_code = f"#{hex_code}"
                        if re.fullmatch(r"#[0-9A-F]{6}", hex_code) and hex_code not in seen:
                            seen.add(hex_code)
                            results.append(hex_code)
                    else:
                        visit(nested)
                return
            if isinstance(value, list):
                for nested in value:
                    visit(nested)
                return
            if isinstance(value, str):
                for match in re.findall(r"#[0-9A-Fa-f]{6}", value):
                    hex_code = match.upper()
                    if hex_code not in seen:
                        seen.add(hex_code)
                        results.append(hex_code)

        visit(brand_context)
        return results

    @classmethod
    def _build_visual_page_review(
        cls,
        *,
        page_index: int,
        image_path: str,
        analysis: dict[str, Any],
        expected_prompt: str,
        page_text: str,
        brand_context: dict[str, Any],
    ) -> dict[str, Any]:
        labels = [
            str(item.get("desc") or "").strip()
            for item in (analysis.get("labels") or [])
            if isinstance(item, dict) and str(item.get("desc") or "").strip()
        ]
        structured_text = [
            entry
            for entry in (analysis.get("structured_text") or [])
            if isinstance(entry, dict)
        ]
        analysis_text = " ".join(
            str(entry.get("text") or entry.get("desc") or "").strip()
            for entry in structured_text
            if str(entry.get("text") or entry.get("desc") or "").strip()
        ).strip()
        observed_text = " ".join(part for part in [str(page_text or "").strip(), analysis_text, " ".join(labels)] if part).strip()
        dominant_colors = [
            str(item.get("hex") or item.get("hex_code") or "").strip().upper()
            for item in (analysis.get("dominant_colors") or [])
            if isinstance(item, dict) and str(item.get("hex") or item.get("hex_code") or "").strip()
        ]
        if not dominant_colors:
            dominant_colors = [
                str(item.get("hex") or item.get("hex_code") or "").strip().upper()
                for item in (analysis.get("text_element_colors") or [])
                if isinstance(item, dict) and str(item.get("hex") or item.get("hex_code") or "").strip()
            ]
        word_count = len(re.findall(r"\b\w+\b", observed_text))
        line_count = max(len([line for line in observed_text.splitlines() if line.strip()]), 1) if observed_text else 0
        box_heights = []
        box_area = 0
        for entry in structured_text:
            bbox = entry.get("bounding_box") if isinstance(entry.get("bounding_box"), dict) else {}
            width = int(bbox.get("w", bbox.get("width", 0)) or 0)
            height = int(bbox.get("h", bbox.get("height", 0)) or 0)
            if width > 0 and height > 0:
                box_heights.append(height)
                box_area += width * height
        canvas_size = cls._image_canvas_size(image_path)
        canvas_width, canvas_height = canvas_size if canvas_size else (0, 0)
        layout_regions = cls._layout_region_diagnostics(
            structured_text=structured_text,
            canvas_width=canvas_width,
            canvas_height=canvas_height,
        )
        hierarchy_signal = 0.0
        if box_heights:
            hierarchy_signal = min(max(max(box_heights) / max(min(box_heights), 1), 1.0), 4.0)
        matched_prompt_terms, missing_prompt_terms = cls._prompt_term_diagnostics(expected_prompt, observed_text, labels)
        prompt_alignment_score = cls._prompt_alignment_score(expected_prompt, observed_text, labels)
        layout_readability_score = cls._layout_readability_score(
            word_count=word_count,
            line_count=line_count,
            text_box_count=len(structured_text),
            hierarchy_signal=hierarchy_signal,
        )
        density_score = cls._density_score(
            word_count=word_count,
            text_box_count=len(structured_text),
            box_area=box_area,
        )
        hierarchy_score = cls._hierarchy_score(
            hierarchy_signal=hierarchy_signal,
            headline_prominence=layout_regions["headline_prominence"],
        )
        crowding_score = cls._crowding_score(
            text_box_count=len(structured_text),
            edge_crowding_count=layout_regions["edge_crowding_count"],
            overlap_count=layout_regions["overlap_count"],
            text_coverage_ratio=layout_regions["text_coverage_ratio"],
        )
        page_balance_score = cls._page_balance_score(layout_regions["vertical_distribution"])
        brand_alignment_score = cls._brand_alignment_score(
            brand_context=brand_context,
            observed_text=observed_text,
            labels=labels,
            dominant_colors=dominant_colors,
        )
        page_findings = cls._page_findings(
            prompt_alignment_score=prompt_alignment_score,
            layout_readability_score=layout_readability_score,
            density_score=density_score,
            hierarchy_score=hierarchy_score,
            crowding_score=crowding_score,
            page_balance_score=page_balance_score,
            missing_prompt_terms=missing_prompt_terms,
        )
        ocr_confidence_score = cls._ocr_confidence_score(
            structured_text=structured_text,
            observed_text=observed_text,
            warning_count=0,
        )
        dominant_region = max(layout_regions["vertical_distribution"].items(), key=lambda item: item[1])[0] if any(layout_regions["vertical_distribution"].values()) else "none"
        return {
            "page_index": page_index,
            "image_path": image_path,
            "ocr_text_excerpt": observed_text[:500],
            "label_hits": labels[:8],
            "text_box_count": len(structured_text),
            "word_count": word_count,
            "line_count": line_count,
            "prompt_alignment_score": prompt_alignment_score,
            "semantic_consistency_score": prompt_alignment_score,
            "layout_readability_score": layout_readability_score,
            "density_score": density_score,
            "hierarchy_score": hierarchy_score,
            "crowding_score": crowding_score,
            "page_balance_score": page_balance_score,
            "ocr_confidence_score": ocr_confidence_score,
            "ocr_confidence_label": cls._ocr_confidence_label(ocr_confidence_score),
            "brand_alignment_score": brand_alignment_score,
            "dominant_colors": dominant_colors[:8],
            "matched_prompt_terms": matched_prompt_terms[:8],
            "missing_prompt_terms": missing_prompt_terms[:8],
            "page_findings": page_findings[:6],
            "region_diagnostics": {
                "dominant_region": dominant_region,
                "edge_crowding_count": layout_regions["edge_crowding_count"],
                "overlap_count": layout_regions["overlap_count"],
                "text_coverage_ratio": round(layout_regions["text_coverage_ratio"], 4),
                "vertical_distribution": layout_regions["vertical_distribution"],
                "headline_prominence": round(layout_regions["headline_prominence"], 2),
            },
            "signals": {
                "hierarchy_signal": round(hierarchy_signal, 2),
                "box_area": box_area,
                "canvas_width": canvas_width,
                "canvas_height": canvas_height,
                "edge_crowding_count": layout_regions["edge_crowding_count"],
                "overlap_count": layout_regions["overlap_count"],
                "text_coverage_ratio": round(layout_regions["text_coverage_ratio"], 4),
                "vertical_distribution": layout_regions["vertical_distribution"],
                "headline_prominence": round(layout_regions["headline_prominence"], 2),
            },
        }

    @staticmethod
    def _ocr_confidence_score(
        *,
        structured_text: list[dict[str, Any]],
        observed_text: str,
        warning_count: int,
    ) -> int:
        confidence_values: list[float] = []
        for entry in structured_text:
            raw = entry.get("confidence", entry.get("score"))
            try:
                value = float(raw)
            except (TypeError, ValueError):
                continue
            if value <= 1.0:
                value *= 100.0
            confidence_values.append(max(0.0, min(100.0, value)))
        if confidence_values:
            base = sum(confidence_values) / max(len(confidence_values), 1)
        else:
            word_count = len(re.findall(r"\b\w+\b", observed_text or ""))
            if word_count >= 40:
                base = 82.0
            elif word_count >= 15:
                base = 72.0
            elif word_count >= 5:
                base = 60.0
            elif observed_text:
                base = 48.0
            else:
                base = 18.0
        base -= warning_count * 6.0
        return max(0, min(100, int(round(base))))

    @staticmethod
    def _ocr_confidence_label(score: int) -> str:
        if score >= 85:
            return "high"
        if score >= 65:
            return "medium"
        if score >= 40:
            return "low"
        return "very_low"

    @staticmethod
    def _document_segments_from_page_reviews(
        page_reviews: list[dict[str, Any]],
        *,
        asset_kind: str,
    ) -> list[dict[str, Any]]:
        if asset_kind not in {"document", "presentation", "image"}:
            return []
        segments: list[dict[str, Any]] = []
        for page_review in page_reviews[:24]:
            excerpt = str(page_review.get("ocr_text_excerpt") or "").strip()
            words = excerpt.split()
            heading_excerpt = " ".join(words[:8]).strip()
            segments.append(
                {
                    "page_index": int(page_review.get("page_index") or 0),
                    "heading_excerpt": heading_excerpt,
                    "text_excerpt": excerpt[:220],
                    "dominant_region": str((page_review.get("region_diagnostics") or {}).get("dominant_region") or "none"),
                    "text_box_count": int(page_review.get("text_box_count") or 0),
                    "word_count": int(page_review.get("word_count") or 0),
                    "ocr_confidence_score": int(page_review.get("ocr_confidence_score") or 0),
                }
            )
        return segments

    @staticmethod
    def _region_overview_from_pages(page_reviews: list[dict[str, Any]]) -> dict[str, Any]:
        overview = {
            "dominant_regions": [],
            "top_region_count": 0,
            "middle_region_count": 0,
            "bottom_region_count": 0,
            "edge_crowding_pages": 0,
            "overlap_pages": 0,
        }
        dominant_regions: list[str] = []
        for page_review in page_reviews:
            region = page_review.get("region_diagnostics") if isinstance(page_review.get("region_diagnostics"), dict) else {}
            dominant = str(region.get("dominant_region") or "").strip()
            if dominant:
                dominant_regions.append(dominant)
            distribution = region.get("vertical_distribution") if isinstance(region.get("vertical_distribution"), dict) else {}
            overview["top_region_count"] += int(distribution.get("top") or 0)
            overview["middle_region_count"] += int(distribution.get("middle") or 0)
            overview["bottom_region_count"] += int(distribution.get("bottom") or 0)
            if int(region.get("edge_crowding_count") or 0) > 0:
                overview["edge_crowding_pages"] += 1
            if int(region.get("overlap_count") or 0) > 0:
                overview["overlap_pages"] += 1
        overview["dominant_regions"] = dominant_regions[:12]
        return overview

    @staticmethod
    def _image_canvas_size(image_path: str) -> tuple[int, int] | None:
        try:
            with Image.open(image_path) as image:
                return int(image.width or 0), int(image.height or 0)
        except Exception:  # noqa: BLE001
            return None

    @classmethod
    def _prompt_term_diagnostics(cls, expected_prompt: str, observed_text: str, labels: list[str]) -> tuple[list[str], list[str]]:
        prompt_tokens = []
        for token in cls._prompt_topic_tokens(expected_prompt):
            if token not in prompt_tokens:
                prompt_tokens.append(token)
        observed_tokens = set(cls._prompt_topic_tokens(f"{observed_text} {' '.join(labels)}", limit=32))
        matched = [token for token in prompt_tokens if token in observed_tokens]
        missing = [token for token in prompt_tokens if token not in observed_tokens]
        return matched, missing

    @staticmethod
    def _layout_region_diagnostics(
        *,
        structured_text: list[dict[str, Any]],
        canvas_width: int,
        canvas_height: int,
    ) -> dict[str, Any]:
        if canvas_width <= 0 or canvas_height <= 0:
            return {
                "edge_crowding_count": 0,
                "overlap_count": 0,
                "text_coverage_ratio": 0.0,
                "vertical_distribution": {"top": 0, "middle": 0, "bottom": 0},
                "headline_prominence": 1.0,
            }
        boxes: list[tuple[int, int, int, int]] = []
        areas: list[int] = []
        distribution = {"top": 0, "middle": 0, "bottom": 0}
        edge_crowding_count = 0
        margin_x = canvas_width * 0.05
        margin_y = canvas_height * 0.05
        for entry in structured_text:
            bbox = entry.get("bounding_box") if isinstance(entry.get("bounding_box"), dict) else {}
            x = int(bbox.get("x", bbox.get("left", 0)) or 0)
            y = int(bbox.get("y", bbox.get("top", 0)) or 0)
            width = int(bbox.get("w", bbox.get("width", 0)) or 0)
            height = int(bbox.get("h", bbox.get("height", 0)) or 0)
            if width <= 0 or height <= 0:
                continue
            boxes.append((x, y, width, height))
            areas.append(width * height)
            center_y = y + (height / 2.0)
            if center_y < canvas_height / 3.0:
                distribution["top"] += 1
            elif center_y < (canvas_height * 2.0) / 3.0:
                distribution["middle"] += 1
            else:
                distribution["bottom"] += 1
            if x <= margin_x or y <= margin_y or (x + width) >= (canvas_width - margin_x) or (y + height) >= (canvas_height - margin_y):
                edge_crowding_count += 1
        overlap_count = 0
        for index, (ax, ay, aw, ah) in enumerate(boxes):
            for bx, by, bw, bh in boxes[index + 1:]:
                overlap_width = min(ax + aw, bx + bw) - max(ax, bx)
                overlap_height = min(ay + ah, by + bh) - max(ay, by)
                if overlap_width > 0 and overlap_height > 0:
                    overlap_count += 1
        total_area = sum(areas)
        coverage_ratio = total_area / max(float(canvas_width * canvas_height), 1.0)
        headline_prominence = 1.0
        if areas:
            sorted_areas = sorted(areas, reverse=True)
            baseline = float(sorted_areas[1] if len(sorted_areas) > 1 else sorted_areas[0] or 1)
            headline_prominence = max(float(sorted_areas[0]) / max(baseline, 1.0), 1.0)
        return {
            "edge_crowding_count": edge_crowding_count,
            "overlap_count": overlap_count,
            "text_coverage_ratio": coverage_ratio,
            "vertical_distribution": distribution,
            "headline_prominence": headline_prominence,
        }

    @staticmethod
    def _hierarchy_score(*, hierarchy_signal: float, headline_prominence: float) -> int:
        score = 100.0
        if hierarchy_signal < 1.2:
            score -= 26.0
        elif hierarchy_signal < 1.5:
            score -= 14.0
        if headline_prominence < 1.15:
            score -= 20.0
        elif headline_prominence < 1.35:
            score -= 10.0
        elif headline_prominence > 2.4:
            score -= 6.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _crowding_score(*, text_box_count: int, edge_crowding_count: int, overlap_count: int, text_coverage_ratio: float) -> int:
        score = 100.0
        if text_box_count > 8:
            score -= min((text_box_count - 8) * 4.0, 18.0)
        score -= min(edge_crowding_count * 6.0, 24.0)
        score -= min(overlap_count * 8.0, 24.0)
        if text_coverage_ratio > 0.35:
            score -= min((text_coverage_ratio - 0.35) * 100.0, 26.0)
        elif text_coverage_ratio < 0.03 and text_box_count > 0:
            score -= 8.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _page_balance_score(vertical_distribution: dict[str, int]) -> int:
        counts = [int(vertical_distribution.get(key) or 0) for key in ("top", "middle", "bottom")]
        total = sum(counts)
        if total <= 1:
            return 100
        mean = total / 3.0
        imbalance = sum(abs(count - mean) for count in counts) / max(total, 1)
        score = 100.0 - min(imbalance * 60.0, 30.0)
        if counts[1] == 0 and counts[0] > 0 and counts[2] > 0:
            score -= 8.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _page_findings(
        *,
        prompt_alignment_score: int,
        layout_readability_score: int,
        density_score: int,
        hierarchy_score: int,
        crowding_score: int,
        page_balance_score: int,
        missing_prompt_terms: list[str],
    ) -> list[str]:
        findings: list[str] = []
        if prompt_alignment_score < 60 and missing_prompt_terms:
            findings.append(f"Important prompt themes seem underrepresented: {', '.join(missing_prompt_terms[:4])}.")
        if layout_readability_score < 60:
            findings.append("Text hierarchy or readability looks weak on this page.")
        if density_score < 60 or crowding_score < 60:
            findings.append("The page appears visually crowded or text-heavy.")
        if hierarchy_score < 60:
            findings.append("Headline prominence is weak relative to the rest of the page.")
        if page_balance_score < 60:
            findings.append("Text distribution across the page feels unbalanced.")
        return findings

    @classmethod
    def _prompt_alignment_score(cls, expected_prompt: str, observed_text: str, labels: list[str]) -> int:
        prompt_tokens = set(cls._prompt_topic_tokens(expected_prompt))
        if not prompt_tokens:
            return 100
        observed_tokens = set(cls._prompt_topic_tokens(f"{observed_text} {' '.join(labels)}", limit=24))
        overlap = len(prompt_tokens & observed_tokens)
        label_overlap_bonus = 1 if any(token in " ".join(labels).lower() for token in prompt_tokens) else 0
        score = ((overlap + label_overlap_bonus) / max(len(prompt_tokens), 1)) * 100.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _layout_readability_score(*, word_count: int, line_count: int, text_box_count: int, hierarchy_signal: float) -> int:
        score = 100.0
        if word_count > 90:
            score -= min((word_count - 90) * 0.7, 35.0)
        if line_count > 8:
            score -= min((line_count - 8) * 4.0, 20.0)
        if text_box_count > 10:
            score -= min((text_box_count - 10) * 3.0, 20.0)
        if hierarchy_signal < 1.2:
            score -= 18.0
        elif hierarchy_signal < 1.6:
            score -= 10.0
        elif hierarchy_signal >= 2.0:
            score += 4.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _density_score(*, word_count: int, text_box_count: int, box_area: int) -> int:
        score = 100.0
        if word_count > 75:
            score -= min((word_count - 75) * 0.9, 40.0)
        if text_box_count > 9:
            score -= min((text_box_count - 9) * 4.0, 24.0)
        if box_area <= 0 and word_count > 0:
            score -= 10.0
        return max(0, min(100, int(round(score))))

    @classmethod
    def _brand_alignment_score(
        cls,
        *,
        brand_context: dict[str, Any],
        observed_text: str,
        labels: list[str],
        dominant_colors: list[str],
    ) -> int:
        score = 70.0
        brand_name_tokens = cls._extract_brand_name_tokens(brand_context)
        observed_tokens = set(cls._prompt_topic_tokens(f"{observed_text} {' '.join(labels)}", limit=24))
        if brand_name_tokens and observed_tokens.intersection(brand_name_tokens):
            score += 15.0
        palette_hexes = set(cls._extract_brand_palette_hexes(brand_context))
        observed_hexes = {
            hex_code if hex_code.startswith("#") else f"#{hex_code}"
            for hex_code in dominant_colors
            if hex_code
        }
        if palette_hexes and observed_hexes:
            overlap = len(palette_hexes & observed_hexes)
            if overlap:
                score += min(overlap * 7.0, 15.0)
            else:
                score -= 12.0
        return max(0, min(100, int(round(score))))

    @staticmethod
    def _average_visual_metric(page_reviews: list[dict[str, Any]], field: str) -> int:
        values = [
            float(page.get(field) or 0.0)
            for page in page_reviews
            if isinstance(page, dict)
        ]
        if not values:
            return 0
        return max(0, min(100, int(round(sum(values) / len(values)))))

    @staticmethod
    def _visual_findings_from_pages(page_reviews: list[dict[str, Any]], *, expected_prompt: str) -> list[str]:
        findings: list[str] = []
        low_alignment = [page["page_index"] for page in page_reviews if int(page.get("prompt_alignment_score") or 0) < 50]
        if low_alignment:
            findings.append(
                f"Prompt/image consistency is weak on page(s) {', '.join(str(item) for item in low_alignment)} compared with the requested topic."
            )
        low_readability = [page["page_index"] for page in page_reviews if int(page.get("layout_readability_score") or 0) < 55]
        if low_readability:
            findings.append(
                f"Layout hierarchy/readability is weak on page(s) {', '.join(str(item) for item in low_readability)}."
            )
        dense_pages = [page["page_index"] for page in page_reviews if int(page.get("density_score") or 0) < 55]
        if dense_pages:
            findings.append(
                f"Text density/clutter appears high on page(s) {', '.join(str(item) for item in dense_pages)}."
            )
        low_brand = [page["page_index"] for page in page_reviews if int(page.get("brand_alignment_score") or 0) < 55]
        if low_brand:
            findings.append(
                f"Visual brand alignment looks weak on page(s) {', '.join(str(item) for item in low_brand)}."
            )
        low_hierarchy = [page["page_index"] for page in page_reviews if int(page.get("hierarchy_score") or 0) < 55]
        if low_hierarchy:
            findings.append(
                f"Headline hierarchy looks weak on page(s) {', '.join(str(item) for item in low_hierarchy)}."
            )
        crowded_pages = [page["page_index"] for page in page_reviews if int(page.get("crowding_score") or 0) < 55]
        if crowded_pages:
            findings.append(
                f"Edge crowding or overlapping text elements appear on page(s) {', '.join(str(item) for item in crowded_pages)}."
            )
        unbalanced_pages = [page["page_index"] for page in page_reviews if int(page.get("page_balance_score") or 0) < 55]
        if unbalanced_pages:
            findings.append(
                f"Page balance feels off on page(s) {', '.join(str(item) for item in unbalanced_pages)}."
            )
        if not findings and expected_prompt:
            findings.append("Visual review did not detect major prompt-alignment, readability, or density issues.")
        return findings[:6]

    def _knowledge_asset_review_text(self, asset) -> str:
        return str(self._knowledge_asset_review_detail(asset).get("content") or "")

    def _knowledge_asset_review_detail(self, asset) -> dict[str, Any]:
        extracted_text = str(getattr(asset, "extracted_text", "") or "").strip()
        if extracted_text:
            return {
                "content": extracted_text[:5000],
                "extraction_method": "extracted_text",
                "gap_note": None,
            }
        extracted_summary = str(getattr(asset, "extracted_summary", "") or "").strip()
        if extracted_summary:
            return {
                "content": extracted_summary,
                "extraction_method": "extracted_summary",
                "gap_note": "Using summary-level extraction instead of full OCR text.",
            }
        normalized = getattr(asset, "normalized_data_json", None)
        if isinstance(normalized, dict) and normalized:
            return {
                "content": "\n".join(
                    f"{key}: {value}"
                    for key, value in normalized.items()
                    if str(value).strip()
                )[:3000],
                "extraction_method": "normalized_data",
                "gap_note": "No extracted text was available; using normalized structured data.",
            }
        structured = getattr(asset, "structured_data_json", None)
        if isinstance(structured, dict) and structured:
            return {
                "content": "\n".join(
                    f"{key}: {value}"
                    for key, value in structured.items()
                    if str(value).strip()
                )[:3000],
                "extraction_method": "structured_data",
                "gap_note": "No extracted text was available; using structured analysis fields.",
            }
        on_demand_content = self._on_demand_asset_text(asset)
        if on_demand_content:
            return {
                "content": on_demand_content[:4000],
                "extraction_method": "on_demand_file_read",
                "gap_note": "Asset was reviewed from a direct file read because no processed extraction was stored.",
            }
        return {
            "content": "",
            "extraction_method": "unavailable",
            "gap_note": "No readable extracted text or structured summary was available for this asset.",
        }

    def _on_demand_asset_text(self, asset) -> str:
        storage_path = str(getattr(asset, "storage_path", "") or "").strip()
        if not storage_path or not self.storage.exists(storage_path):
            return ""
        absolute_path = self.storage.absolute_path(storage_path)
        suffix = Path(str(getattr(asset, "original_filename", "") or absolute_path)).suffix.lower()
        try:
            if suffix in {".txt", ".md", ".csv"}:
                return Path(absolute_path).read_text(encoding="utf-8", errors="replace").strip()
            if suffix == ".json":
                raw = Path(absolute_path).read_text(encoding="utf-8", errors="replace").strip()
                if not raw:
                    return ""
                try:
                    parsed = json.loads(raw)
                except json.JSONDecodeError:
                    return raw
                return json.dumps(parsed, ensure_ascii=False, indent=2)
            if suffix == ".docx":
                doc = Document(absolute_path)
                return "\n".join(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()).strip()
            if suffix in {".png", ".jpg", ".jpeg", ".webp", ".pdf"}:
                extracted = self.ocr.extract(absolute_path)
                text = str(extracted.get("text") or "").strip()
                if text:
                    return text
                analysis_path = extracted.get("analysis_path")
                if analysis_path and Path(str(analysis_path)).exists():
                    return Path(str(analysis_path)).read_text(encoding="utf-8", errors="replace").strip()
        except Exception:
            return ""
        return ""

    @staticmethod
    def _asset_kind_from_mime(mime_type: str) -> str:
        lowered = str(mime_type or "").lower()
        if lowered.startswith("image/"):
            return "image"
        if lowered == "application/pdf":
            return "document"
        if "word" in lowered or "document" in lowered or lowered.endswith("docx"):
            return "document"
        if "presentation" in lowered or "powerpoint" in lowered:
            return "presentation"
        if lowered in {"text/plain", "text/markdown", "application/json", "text/csv"}:
            return "text"
        return "asset"

    def _knowledge_asset_kind(self, asset) -> str:
        mime_type = str(getattr(asset, "mime_type", "") or "").strip()
        kind = self._asset_kind_from_mime(mime_type)
        if kind != "asset":
            return kind
        suffix = Path(str(getattr(asset, "original_filename", "") or "")).suffix.lower()
        if suffix in {".pdf", ".docx", ".doc", ".pptx", ".ppt"}:
            return "document"
        if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            return "image"
        return "asset"

    @staticmethod
    def _quality_value(*sources: Any, key: str, default: float = 0.0) -> float:
        for source in sources:
            if not isinstance(source, dict):
                continue
            quality = source.get("analysis_quality") if isinstance(source.get("analysis_quality"), dict) else source
            try:
                if key in quality:
                    return float(quality.get(key) or 0.0)
            except (TypeError, ValueError):
                continue
        return default

    @staticmethod
    def _quality_list(*sources: Any, key: str) -> list[str]:
        items: list[str] = []
        for source in sources:
            if not isinstance(source, dict):
                continue
            quality = source.get("analysis_quality") if isinstance(source.get("analysis_quality"), dict) else source
            values = quality.get(key)
            if not isinstance(values, list):
                continue
            for item in values:
                text = str(item or "").strip()
                if text and text not in items:
                    items.append(text)
        return items

    def _knowledge_asset_diagnostics(self, asset) -> dict[str, Any]:
        metadata = getattr(asset, "metadata_json", None)
        normalized = getattr(asset, "normalized_data_json", None)
        structured = getattr(asset, "structured_data_json", None)
        validation_summary = getattr(asset, "validation_summary_json", None)
        warnings = []
        if isinstance(validation_summary, dict):
            warnings = [str(item).strip() for item in (validation_summary.get("warnings") or []) if str(item).strip()]
        diagnostics = {
            "source_format": str((metadata or {}).get("source_format") or Path(str(getattr(asset, "original_filename", "") or "")).suffix.lower().lstrip(".") or "").strip() or None,
            "validation_state": str(getattr(asset, "validation_state", "") or "").strip() or None,
            "warning_count": len(warnings),
            "warnings": warnings[:5],
            "analysis_quality_score": self._quality_value(metadata, normalized, structured, key="analysis_quality_score"),
            "summary_quality_score": self._quality_value(metadata, normalized, structured, key="summary_quality_score"),
            "ocr_signal_score": self._quality_value(metadata, normalized, structured, key="ocr_signal_score"),
            "ocr_noise_ratio": self._quality_value(metadata, normalized, structured, key="ocr_noise_ratio"),
            "promotional_line_ratio": self._quality_value(metadata, normalized, structured, key="promotional_line_ratio"),
            "evidence_types": self._quality_list(metadata, normalized, structured, key="evidence_types"),
            "observed_signal_types": self._quality_list(metadata, normalized, structured, key="observed_signal_types"),
            "available_signal_types": self._quality_list(metadata, normalized, structured, key="available_signal_types"),
        }
        diagnostics["visual_signal_count"] = len(diagnostics["observed_signal_types"]) or len(diagnostics["available_signal_types"])
        return diagnostics

    @classmethod
    def _asset_diagnostics(cls, asset_review_blocks: list[dict[str, Any]]) -> dict[str, Any]:
        if not asset_review_blocks:
            return {
                "asset_types": [],
                "source_types": [],
                "document_count": 0,
                "image_count": 0,
                "generated_count": 0,
                "uploaded_count": 0,
                "source_quality_score": 100,
                "document_structure_score": 100,
                "visual_diagnostic_score": 100,
                "prompt_alignment_score": 100,
                "layout_readability_score": 100,
                "density_score": 100,
                "brand_alignment_score": 100,
                "hierarchy_score": 100,
                "crowding_score": 100,
                "page_balance_score": 100,
                "ocr_confidence_score": 100,
            }
        asset_types = list(dict.fromkeys(str(block.get("asset_kind") or "asset") for block in asset_review_blocks))
        source_types = list(dict.fromkeys(str(block.get("source_type") or "asset") for block in asset_review_blocks))
        document_blocks = [block for block in asset_review_blocks if block.get("asset_kind") in {"document", "presentation"}]
        image_blocks = [block for block in asset_review_blocks if block.get("asset_kind") == "image"]
        warning_penalty = sum(int((block.get("diagnostics") or {}).get("warning_count") or 0) for block in asset_review_blocks) * 8
        source_quality_components: list[float] = []
        visual_components: list[float] = []
        document_components: list[float] = []
        prompt_alignment_components: list[float] = []
        readability_components: list[float] = []
        density_components: list[float] = []
        brand_alignment_components: list[float] = []
        hierarchy_components: list[float] = []
        crowding_components: list[float] = []
        balance_components: list[float] = []
        ocr_confidence_components: list[float] = []
        for block in asset_review_blocks:
            diagnostics = block.get("diagnostics") if isinstance(block.get("diagnostics"), dict) else {}
            analysis_quality_score = float(diagnostics.get("analysis_quality_score") or 0.0)
            summary_quality_score = float(diagnostics.get("summary_quality_score") or 0.0)
            ocr_signal_score = float(diagnostics.get("ocr_signal_score") or 0.0)
            ocr_noise_ratio = float(diagnostics.get("ocr_noise_ratio") or 0.0)
            visual_signal_count = int(diagnostics.get("visual_signal_count") or 0)
            prompt_alignment_components.append(float(diagnostics.get("prompt_alignment_score") or 0.0))
            readability_components.append(float(diagnostics.get("layout_readability_score") or 0.0))
            density_components.append(float(diagnostics.get("density_score") or 0.0))
            brand_alignment_components.append(float(diagnostics.get("brand_alignment_score") or 0.0))
            hierarchy_components.append(float(diagnostics.get("hierarchy_score") or 0.0))
            crowding_components.append(float(diagnostics.get("crowding_score") or 0.0))
            balance_components.append(float(diagnostics.get("page_balance_score") or 0.0))
            ocr_confidence_components.append(float(diagnostics.get("ocr_confidence_score") or 0.0))
            source_quality_components.append(min(100.0, max(0.0, (analysis_quality_score * 8.0) + (summary_quality_score * 2.0) + (ocr_signal_score * 2.0) - (ocr_noise_ratio * 20.0))))
            if block.get("asset_kind") == "image":
                visual_components.append(
                    min(
                        100.0,
                        max(
                            0.0,
                            (analysis_quality_score * 6.0)
                            + (visual_signal_count * 5.0)
                            + (float(diagnostics.get("prompt_alignment_score") or 0.0) * 0.25)
                            + (float(diagnostics.get("layout_readability_score") or 0.0) * 0.2)
                            + (float(diagnostics.get("density_score") or 0.0) * 0.15)
                            + (float(diagnostics.get("brand_alignment_score") or 0.0) * 0.15)
                            + (float(diagnostics.get("hierarchy_score") or 0.0) * 0.15)
                            + (float(diagnostics.get("crowding_score") or 0.0) * 0.1)
                            + (float(diagnostics.get("page_balance_score") or 0.0) * 0.1)
                            - (ocr_noise_ratio * 20.0)
                        ),
                    )
                )
            if block.get("asset_kind") in {"document", "presentation"}:
                page_count = int(block.get("page_count") or 0)
                extraction_method = str(block.get("extraction_method") or "")
                extraction_bonus = 18.0 if extraction_method in {"extracted_text", "on_demand_file_read"} else 8.0
                document_components.append(min(100.0, max(0.0, (summary_quality_score * 9.0) + (ocr_signal_score * 3.0) + extraction_bonus + min(page_count, 8) * 2.0 - (ocr_noise_ratio * 20.0))))
        source_quality_score = max(0, min(100, int(round((sum(source_quality_components) / max(len(source_quality_components), 1)) - warning_penalty))))
        document_structure_score = max(0, min(100, int(round(sum(document_components) / max(len(document_components), 1))))) if document_components else 100
        visual_diagnostic_score = max(0, min(100, int(round(sum(visual_components) / max(len(visual_components), 1))))) if visual_components else 100
        prompt_alignment_score = max(0, min(100, int(round(sum(prompt_alignment_components) / max(len(prompt_alignment_components), 1))))) if prompt_alignment_components else 100
        layout_readability_score = max(0, min(100, int(round(sum(readability_components) / max(len(readability_components), 1))))) if readability_components else 100
        density_score = max(0, min(100, int(round(sum(density_components) / max(len(density_components), 1))))) if density_components else 100
        brand_alignment_score = max(0, min(100, int(round(sum(brand_alignment_components) / max(len(brand_alignment_components), 1))))) if brand_alignment_components else 100
        hierarchy_score = max(0, min(100, int(round(sum(hierarchy_components) / max(len(hierarchy_components), 1))))) if hierarchy_components else 100
        crowding_score = max(0, min(100, int(round(sum(crowding_components) / max(len(crowding_components), 1))))) if crowding_components else 100
        page_balance_score = max(0, min(100, int(round(sum(balance_components) / max(len(balance_components), 1))))) if balance_components else 100
        ocr_confidence_score = max(0, min(100, int(round(sum(ocr_confidence_components) / max(len(ocr_confidence_components), 1))))) if ocr_confidence_components else 100
        return {
            "asset_types": asset_types,
            "source_types": source_types,
            "document_count": len(document_blocks),
            "image_count": len(image_blocks),
            "generated_count": sum(1 for block in asset_review_blocks if block.get("source_type") == "generated_asset"),
            "uploaded_count": sum(1 for block in asset_review_blocks if block.get("source_type") == "knowledge_asset"),
            "source_quality_score": source_quality_score,
            "document_structure_score": document_structure_score,
            "visual_diagnostic_score": visual_diagnostic_score,
            "prompt_alignment_score": prompt_alignment_score,
            "layout_readability_score": layout_readability_score,
            "density_score": density_score,
            "brand_alignment_score": brand_alignment_score,
            "hierarchy_score": hierarchy_score,
            "crowding_score": crowding_score,
            "page_balance_score": page_balance_score,
            "ocr_confidence_score": ocr_confidence_score,
        }

    @staticmethod
    def _visual_review_report(asset_review_blocks: list[dict[str, Any]]) -> dict[str, Any]:
        visual_blocks = [
            block
            for block in asset_review_blocks
            if isinstance(block.get("visual_review"), dict) and block.get("visual_review")
        ]
        if not visual_blocks:
            return {
                "asset_count": 0,
                "page_count": 0,
                "prompt_alignment_score": 100,
                "layout_readability_score": 100,
                "density_score": 100,
                "brand_alignment_score": 100,
                "hierarchy_score": 100,
                "crowding_score": 100,
                "page_balance_score": 100,
                "ocr_confidence_score": 100,
                "visual_diagnostic_score": 100,
                "findings": [],
                "page_reviews": [],
                "document_segments": [],
                "region_overview": {},
            }
        page_reviews = [
            {
                "asset_id": block.get("asset_id"),
                "asset_name": block.get("asset_name"),
                **page_review,
            }
            for block in visual_blocks
            for page_review in (block.get("visual_review", {}).get("page_reviews") or [])
            if isinstance(page_review, dict)
        ]
        findings = []
        for block in visual_blocks:
            for finding in (block.get("visual_review", {}).get("findings") or []):
                text = str(finding or "").strip()
                if text and text not in findings:
                    findings.append(text)
        return {
            "asset_count": len(visual_blocks),
            "page_count": len(page_reviews),
            "prompt_alignment_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("prompt_alignment_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "layout_readability_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("layout_readability_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "density_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("density_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "brand_alignment_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("brand_alignment_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "hierarchy_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("hierarchy_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "crowding_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("crowding_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "page_balance_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("page_balance_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "ocr_confidence_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("ocr_confidence_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "visual_diagnostic_score": max(0, min(100, int(round(sum(float((block.get("visual_review", {}).get("diagnostics") or {}).get("visual_diagnostic_score") or 0.0) for block in visual_blocks) / max(len(visual_blocks), 1))))),
            "findings": findings[:8],
            "page_reviews": page_reviews[:20],
            "document_segments": [
                {
                    "asset_id": block.get("asset_id"),
                    "asset_name": block.get("asset_name"),
                    **segment,
                }
                for block in visual_blocks
                for segment in (block.get("visual_review", {}).get("document_segments") or [])
                if isinstance(segment, dict)
            ][:24],
            "region_overview": {
                "dominant_regions": [
                    str(region).strip()
                    for block in visual_blocks
                    for region in ((block.get("visual_review", {}).get("region_overview") or {}).get("dominant_regions") or [])
                    if str(region).strip()
                ][:16],
                "edge_crowding_pages": sum(
                    int(((block.get("visual_review", {}).get("region_overview") or {}).get("edge_crowding_pages") or 0))
                    for block in visual_blocks
                ),
                "overlap_pages": sum(
                    int(((block.get("visual_review", {}).get("region_overview") or {}).get("overlap_pages") or 0))
                    for block in visual_blocks
                ),
            },
        }

    @staticmethod
    def _asset_coverage_score(asset_review_blocks: list[dict[str, Any]]) -> int:
        if not asset_review_blocks:
            return 100
        strong_methods = {"generated_payload", "extracted_text", "on_demand_file_read"}
        total = len(asset_review_blocks)
        strong = sum(1 for block in asset_review_blocks if str(block.get("extraction_method") or "") in strong_methods)
        ratio = strong / max(total, 1)
        return max(0, min(100, int(round(ratio * 100))))

    @staticmethod
    def _source_observations(asset_review_blocks: list[dict[str, Any]]) -> list[str]:
        observations: list[str] = []
        if not asset_review_blocks:
            return observations
        generated_count = sum(1 for block in asset_review_blocks if block.get("source_type") == "generated_asset")
        uploaded_count = sum(1 for block in asset_review_blocks if block.get("source_type") == "knowledge_asset")
        if generated_count:
            observations.append(f"Reviewed {generated_count} generated asset(s) linked to previous content.")
        if uploaded_count:
            observations.append(f"Reviewed {uploaded_count} uploaded knowledge asset(s) using visual-first ingestion plus extracted or structured text.")
        fallback_count = sum(
            1
            for block in asset_review_blocks
            if str(block.get("extraction_method") or "") in {"extracted_summary", "normalized_data", "structured_data", "on_demand_file_read"}
        )
        if fallback_count:
            observations.append(f"{fallback_count} asset(s) relied on fallback extraction methods instead of a full stored OCR body.")
        visual_count = sum(1 for block in asset_review_blocks if isinstance(block.get("visual_review"), dict) and block.get("visual_review"))
        if visual_count:
            observations.append(f"Ran OCR-first visual diagnostics for {visual_count} image/document asset(s).")
        visual_first_count = sum(1 for block in asset_review_blocks if str(block.get("review_workflow") or "") == "visual_first")
        if visual_first_count:
            observations.append(f"{visual_first_count} asset(s) were treated as first-class visual review inputs rather than text-only artifacts.")
        return observations

    @staticmethod
    def _evaluation_summary(
        report: dict[str, Any],
        *,
        asset_review_blocks: list[dict[str, Any]] | None = None,
        asset_coverage_score: int | None = None,
    ) -> str:
        score = int(report.get("score") or 0)
        matched = report.get("matched_signals") if isinstance(report.get("matched_signals"), list) else []
        deviations = report.get("deviations") if isinstance(report.get("deviations"), list) else []
        strengths = ", ".join(str(item) for item in matched[:2]) or "no strong aligned signals yet"
        risks = ", ".join(str(item) for item in deviations[:2]) or "no major issues detected"
        if asset_review_blocks:
            return (
                f"Tone score: {score}/100. Asset-backed review coverage: {int(asset_coverage_score or 0)}/100. "
                f"Strongest signals: {strengths}. Biggest gaps: {risks}."
            )
        return f"Tone score: {score}/100. Strongest signals: {strengths}. Biggest gaps: {risks}."
